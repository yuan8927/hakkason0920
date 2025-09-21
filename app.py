# -*- coding: utf-8 -*-
import hashlib
from datetime import datetime
from typing import List

import numpy as np
import pandas as pd
import pytz
import streamlit as st

import gspread
from google.oauth2.service_account import Credentials
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side
from openpyxl.utils import get_column_letter

# =============== 設定 ===============
JST = pytz.timezone("Asia/Tokyo")

# Secrets 読み込み（Streamlit Cloud の Secrets から）
APP_SECRETS = st.secrets.get("app", {})
ALLOWED_DATES: List[str] = list(APP_SECRETS.get("allowed_dates", ["2025-10-25"]))
ALLOWED_PLACES: List[str] = list(APP_SECRETS.get("allowed_places", ["メインステージ"]))
DAY_START = APP_SECRETS.get("day_start", "09:00")
DAY_END = APP_SECRETS.get("day_end", "18:00")
GSHEET_ID = APP_SECRETS.get("gsheet_id", "")

# =============== Google Sheets 接続 ===============
@st.cache_resource(show_spinner=False)
def get_worksheet():
    creds = Credentials.from_service_account_info(st.secrets["gcp_service_account"])
    scoped = creds.with_scopes([
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive"
    ])
    gc = gspread.authorize(scoped)

    try:
        sh = gc.open_by_key(GSHEET_ID)
    except Exception as ex:
        st.error(f"スプレッドシートを開けませんでした。GSHEET_ID={GSHEET_ID}, Error={ex}")
        raise

    try:
        ws = sh.worksheet("data")
    except gspread.WorksheetNotFound:
        ws = sh.add_worksheet(title="data", rows=1000, cols=10)
        ws.append_row(["timestamp", "user_name", "date", "place", "start", "end", "priority"])
    return ws


# =============== ユーティリティ ===============
def time_slots(day_start: str, day_end: str, step_min: int = 15) -> List[str]:
    base = pd.to_datetime(f"2000-01-01 {day_start}")
    end = pd.to_datetime(f"2000-01-01 {day_end}")
    # 端点ともに含む（例: 09:00, 09:15, ..., 18:00）
    return pd.date_range(base, end, freq=f"{step_min}min").strftime("%H:%M").tolist()

SLOTS = time_slots(DAY_START, DAY_END, 15)

def validate_range(start: str, end: str) -> bool:
    return pd.to_datetime(start) < pd.to_datetime(end)

def name_to_color(name: str) -> str:
    """利用者名から安定した淡色 HEX を生成。"""
    palette = [
        "#CFE8FF", "#FFD6A5", "#B9FBC0", "#FFADAD", "#FDFFB6", "#A0C4FF",
        "#CAFFBF", "#9BF6FF", "#F1C0E8", "#BDB2FF", "#FFC6FF", "#E0F7FA",
    ]
    h = int(hashlib.md5(name.encode("utf-8")).hexdigest(), 16)
    return palette[h % len(palette)]

def append_rows(ws, rows: list[list[str]]):
    ws.append_rows(rows, value_input_option="USER_ENTERED")


@st.cache_data(ttl=30)
def load_df() -> pd.DataFrame:   ###多少変更した
    ws = get_worksheet()  # ← ここで取得

    records = ws.get_all_records()
    df = pd.DataFrame(records)
    if df.empty:
        df = pd.DataFrame(columns=["timestamp", "user_name", "date", "place", "start", "end", "priority", "remarks"])  # remarks追加
    for c in ["date", "start", "end"]:
        if c in df.columns:
            df[c] = df[c].astype(str)
    if "priority" in df.columns:
        df["priority"] = pd.to_numeric(df["priority"], errors="coerce").astype("Int64")
    return df


def make_excel_by_date(df: pd.DataFrame, date_str: str) -> str:
    """指定日のデータから、場所ごとにシートを作る Excel を生成し、ファイルパスを返す。"""
    df_day = df[df["date"] == date_str].copy()
    if df_day.empty:
        raise ValueError("この日付のデータがありません。")

    wb = Workbook()
    wb.remove(wb.active)

    border_thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    for place in ALLOWED_PLACES:
        df_p = df_day[df_day["place"] == place].copy()
        ws = wb.create_sheet(title=place[:31])  # シート名は31文字制限

        # ヘッダー（A1=利用者, B1〜=時間スロット）
        header = ["利用者"] + SLOTS
        ws.append(header)
        for col_idx in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 対象利用者（この場所に希望を出した人）
        users = df_p["user_name"].dropna().unique().tolist()

        for r, user in enumerate(users, start=2):
            ws.cell(row=r, column=1, value=user)
            ws.cell(row=r, column=1).alignment = Alignment(vertical="center")

            sub = df_p[df_p["user_name"] == user]
            user_color = name_to_color(user)
            fill = PatternFill(start_color=user_color.replace('#',''), end_color=user_color.replace('#',''), fill_type="solid")

            for _, rec in sub.iterrows():
                start, end, pr = str(rec["start"]), str(rec["end"]), int(rec["priority"])
                
                try:
                    start = pd.to_datetime(start).strftime("%H:%M")
                    end = pd.to_datetime(end).strftime("%H:%M")
                except Exception:
                    # 時刻変換できないものはスキップ
                    continue
                
                if not validate_range(start, end):
                    continue
                # 開始・終了のスロット index（終了は“含めない”開区間）
                try:
                    s_idx = SLOTS.index(start)
                    e_idx = SLOTS.index(end)
                except ValueError:
                    print(f"[WARN] Slot not found: start={start}, end={end}, SLOTS[0]={SLOTS[0]}")
                    continue
                
                # Excel の列番号（A=1, B=2 ...）: B 列が SLOTS[0]
                start_col = 2 + s_idx
                end_col_exclusive = 2 + e_idx  # ここは含めない終端

                for c in range(start_col, end_col_exclusive):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = fill
                    label = f"第{pr}希望"
                    cell.value = f"{cell.value},{label}" if cell.value else label
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # 罫線・列幅
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = border_thin
        ws.column_dimensions['A'].width = 18
        for col in range(2, len(SLOTS) + 2):
            col_letter = get_column_letter(col)
            ws.column_dimensions[col_letter].width = 4.2

    out_name = f"{date_str.replace('-', '')}.xlsx"
    wb.save(out_name)
    return out_name

# =============== UI ===============
st.set_page_config(page_title="施設利用希望フォーム", layout="wide")
st.title("施設利用希望 収集・管理アプリ")

ws = get_worksheet()

user_tab, admin_tab = st.tabs(["📝 利用者フォーム", "🛠 管理（一覧・Excel出力）"])

with user_tab:
    st.caption("※ 第1〜第3希望はすべて必須です。時間は15分刻みで選択してください。")

    name = st.text_input("お名前（必須）")

    remarks = st.text_area(
    "希望理由・備考（任意）",
    placeholder="希望理由や備考があれば入力してください",
    height=120
    )

    def hope_block(title: str):
        st.subheader(title)
        c1, c2, c3, c4 = st.columns([1.2, 1.2, 1, 1])
        with c1:
            sel_date = st.selectbox("日付", ALLOWED_DATES, key=f"date_{title}")
        with c2:
            sel_place = st.selectbox("場所", ALLOWED_PLACES, key=f"place_{title}")
        with c3:
            start = st.selectbox("開始", SLOTS[:-1], key=f"start_{title}")
        with c4:
            end = st.selectbox("終了", SLOTS[1:], key=f"end_{title}")
        return sel_date, sel_place, start, end

    d1, p1, s1, e1 = hope_block("第1希望")
    d2, p2, s2, e2 = hope_block("第2希望")
    d3, p3, s3, e3 = hope_block("第3希望")

    if st.button("送信する", type="primary"):
        errors = []

        name_input = name.strip()

        if not name_input:
            errors.append("お名前は必須です。")
        else:
            ###多少変更した
            #名前の重複チェック
            #  ※ 正規化して比較（前後/連続スペース、全角→半角スペース、大小文字差を吸収）
            def normalize_name(s:str) -> str:
                s = str(s).strip().replace("　"," ")
                s = " ".join(s.split())
                return s.lower()

            existing_names = [r.get("user_name","") for r in ws.get_all_records()]
            existing_norm = {normalize_name(n) for n in existing_names}
            if normalize_name(name_input) in existing_norm:
                errors.append(f"この名前「{name_input}」は既に登録されています。別の名前を入力してください。")

        for idx, (s, e) in enumerate([(s1, e1), (s2, e2), (s3, e3)], start=1):
            if not validate_range(s, e):
                errors.append(f"第{idx}希望の時間範囲が不正です（開始 < 終了）。")
        if errors:
            st.error("\n".join(errors))
        else:
            ts = datetime.now(JST).strftime("%Y-%m-%d %H:%M:%S")
            rows = [
                [ts, name_input, d1, p1, s1, e1, 1, remarks],
                [ts, name_input, d2, p2, s2, e2, 2, remarks],
                [ts, name_input, d3, p3, s3, e3, 3, remarks],
            ]
            try:
                append_rows(ws, rows)
                st.success("送信しました。ご協力ありがとうございます！")
                load_df.clear()  # キャッシュ削除
            except Exception as ex:
                st.error(f"送信に失敗しました: {ex}")

with admin_tab:
    st.subheader("データ一覧（最新）")
    df = load_df()
    st.dataframe(df, use_container_width=True)

    st.divider()
    st.subheader("Excel 出力（ガントチャート風）")
    selectable_dates = sorted(df["date"].dropna().unique().tolist()) if not df.empty else []
    target_dates = st.multiselect("作成する日付を選択", options=selectable_dates, default=selectable_dates)

    if st.button("選択した日付のExcelを作成"):
        if not target_dates:
            st.warning("対象日付がありません。")
        else:
            for d in target_dates:
                try:
                    path = make_excel_by_date(df, d)
                    with open(path, "rb") as f:
                        st.download_button(
                            label=f"{d} のExcelをダウンロード",
                            file_name=path,
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            data=f.read(),
                        )
                except Exception as ex:
                    st.error(f"{d} の生成に失敗: {ex}")
                    
with st.tab("確認＆キャンセル"):
    st.subheader("予約確認・キャンセル")

    search_name = st.text_input("名前で検索", "")

    if search_name:
        df = load_df()
        # 名前正規化して検索
        def normalize_name(s):
            return str(s).strip().replace("　"," ").lower()
        
        target_df = df[df["user_name"].apply(lambda x: normalize_name(x) == normalize_name(search_name))]
        
        if target_df.empty:
            st.info("該当する予約はありません。")
        else:
            st.dataframe(target_df[["date","place","start","end","priority","remarks"]], use_container_width=True)

            if st.button(f"{search_name} の予約をキャンセル"):
                ws = get_worksheet()
                # Google Sheets 側で該当行を消す
                all_records = ws.get_all_records()
                # 削除対象行番号を取得
                row_indices_to_delete = [
                    i+2 for i, r in enumerate(all_records) 
                    if normalize_name(r.get("user_name","")) == normalize_name(search_name)
                ]
                for row_idx in reversed(row_indices_to_delete):  # 下から削除すると行番号ずれない
                    ws.delete_row(row_idx)
                st.success(f"{search_name} の予約を削除しました。")
                load_df.clear()  # キャッシュクリア